"""lib/excel.py : édition de l'Excel maître sur le vrai template du projet."""
from __future__ import annotations

import openpyxl
import pytest

from excel_templates import write_starter_excel
from lib import excel as XL


@pytest.fixture
def classeur(project):
    chemin = project / f"{project.name}_sessions.xlsx"
    write_starter_excel(chemin, "single", project.name)
    return chemin


def test_lire_feuille_en_texte(classeur):
    df = XL.lire_feuille(classeur, "Sessions")
    assert list(df["id"]) == ["971", "970-M1", "970-M2"]
    assert list(df["mouse_id"]) == ["971", "970", "970"]
    assert XL.feuilles_editables(classeur) == ["Sessions"]


def test_enregistrer_modifie_ajoute_supprime_sans_casser_le_classeur(project, classeur):
    avant = XL.lire_feuille(classeur, "Sessions")
    apres = avant.copy()
    apres.at[0, "group"] = "MCCf/f"            # modification
    apres = apres.drop(index=1)                 # suppression de 970-M1
    apres.loc[3] = [None] * len(apres.columns)  # ligne ajoutée (index >= len)
    apres.at[3, "id"] = "972"
    apres.at[3, "mouse_id"] = "972"
    apres = apres.assign(regime="")             # colonne ajoutée
    apres.at[0, "regime"] = "HFD"

    sauvegarde = XL.enregistrer(project, classeur, "Sessions", avant, apres)

    assert sauvegarde.is_file() and sauvegarde.parent == XL.dossier_sauvegardes(project)
    ws = openpyxl.load_workbook(classeur)["Sessions"]
    lignes = list(ws.iter_rows(values_only=True))
    entetes = lignes[0]
    assert entetes[-1] == "regime"
    ids = [r[0] for r in lignes[1:] if r[0] is not None]
    assert ids == ["971", "970-M2", "972"]
    assert lignes[1][entetes.index("group")] == "MCCf/f"
    assert lignes[1][entetes.index("regime")] == "HFD"
    # Types : `id` texte reste texte, `mouse_id` numérique reste nombre,
    # et une nouvelle cellule numérique devient un nombre.
    assert lignes[1][entetes.index("mouse_id")] == 971
    assert lignes[3][entetes.index("mouse_id")] == 972
    # Mise en forme du template intacte.
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert openpyxl.load_workbook(classeur).sheetnames[0] == "Instructions"


def test_id_en_double_refuse_et_fichier_intact(project, classeur):
    original = classeur.read_bytes()
    avant = XL.lire_feuille(classeur, "Sessions")
    apres = avant.copy()
    apres.at[1, "id"] = "971"
    with pytest.raises(ValueError, match="doublons"):
        XL.enregistrer(project, classeur, "Sessions", avant, apres)
    assert classeur.read_bytes() == original
    assert not (XL.dossier_sauvegardes(project) / "en-cours.xlsx").exists()


def test_a_change(classeur):
    avant = XL.lire_feuille(classeur, "Sessions")
    apres = avant.copy()
    assert not XL.a_change(avant, apres)
    apres.at[0, "notes"] = None   # cellule vide côté éditeur == "" côté fichier
    apres.at[0, "notes"] = avant.at[0, "notes"]
    assert not XL.a_change(avant, apres)
    apres.at[0, "notes"] = "x"
    assert XL.a_change(avant, apres)


def test_valider_refuse_un_fichier_qui_n_est_pas_un_classeur(tmp_path):
    faux = tmp_path / "faux.xlsx"
    faux.write_bytes(b"pas un xlsx")
    with pytest.raises(ValueError):
        XL.valider(faux)
