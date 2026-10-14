"""Lire et modifier l'Excel maître depuis l'app — l'Excel reste la source de vérité.

Écriture cellule par cellule avec openpyxl, jamais `DataFrame.to_excel` :
seules les cellules modifiées sont réécrites, donc les styles du template
(en-tête bleu, lignes d'exemple grisées), le volet figé, les largeurs de
colonnes, les autres feuilles et les formules des cellules non touchées
survivent. Le classeur est rechargé au moment d'enregistrer : une
modification faite entre-temps dans Excel sur une autre cellule n'est pas
écrasée.

Avant de remplacer le fichier, le nouveau classeur est relu par
`sync_from_excel.py` lui-même (`detect_schema`, `parse_single`) : un `id`
en double ou une feuille cassée est refusé ici plutôt qu'au sync suivant.
L'ancienne version est copiée dans `<projet>/.backups/excel/`.

Ce module n'importe pas Streamlit.
"""
from __future__ import annotations

import datetime as dt
import os
import shutil
from pathlib import Path

import openpyxl
import pandas as pd

from lib.project import SCRIPTS_DIR  # noqa: F401 — met scripts/ dans sys.path
from paths import raw_dir  # noqa: E402
from sync_from_excel import detect_schema, parse_single  # noqa: E402

# Feuilles de données des deux schémas ; `Instructions` et un éventuel
# `Codebook` ne sont pas proposés à l'édition.
FEUILLES = ("Sessions", "Subjects", "Trials_Videos", "Arena_Mapping")


class ExcelVerrouille(Exception):
    """Le fichier est ouvert ailleurs (Excel sous Windows le verrouille)."""


def dossier_sauvegardes(projet: Path) -> Path:
    return Path(projet) / ".backups" / "excel"


def nom_sauvegarde(chemin: Path) -> str:
    return f"{chemin.stem}_{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx"


def ouvert_dans_excel(chemin: Path) -> bool:
    """Excel (Windows et macOS) pose un `~$<nom>` à côté du fichier ouvert."""
    return (chemin.parent / f"~${chemin.name}").exists()


def feuilles_editables(chemin: Path) -> list[str]:
    wb = openpyxl.load_workbook(chemin, read_only=True)
    try:
        return [f for f in FEUILLES if f in wb.sheetnames]
    finally:
        wb.close()


def _texte(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime) and v.time() == dt.time():
        return v.date().isoformat()
    if isinstance(v, (dt.date, dt.datetime, dt.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _norm(v) -> str:
    """Valeur renvoyée par `st.data_editor` -> texte comparable."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def lire_feuille(chemin: Path, feuille: str) -> pd.DataFrame:
    """La feuille en texte, ligne `i` du DataFrame = ligne Excel `i + 2`.

    Tout en texte : une colonne qui mêle `971` et `970-M1` ferait échouer
    la sérialisation Arrow de `st.data_editor`. Les types sont rétablis à
    l'écriture (`_valeur`). Seules les lignes et colonnes vides en FIN de
    feuille sont retirées, pour garder la correspondance de positions.
    """
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    try:
        lignes = [[_texte(v) for v in r]
                  for r in wb[feuille].iter_rows(values_only=True)]
    finally:
        wb.close()
    if not lignes:
        return pd.DataFrame()
    while len(lignes) > 1 and not any(lignes[-1]):
        lignes.pop()
    largeur = max(len(r) for r in lignes)
    lignes = [r + [""] * (largeur - len(r)) for r in lignes]
    while largeur and not any(r[largeur - 1] for r in lignes):
        largeur -= 1
    entetes = [h or f"(colonne {j + 1})" for j, h in enumerate(lignes[0][:largeur])]
    return pd.DataFrame([r[:largeur] for r in lignes[1:]], columns=entetes)


def _valeur(texte, ancienne):
    """Texte saisi -> valeur de cellule, dans le type de l'ancienne si possible.

    `"971"` saisi dans une cellule texte reste du texte (sync_from_excel
    lit `id` tel quel) ; dans une cellule numérique ou vide, redevient un
    nombre. Une saisie illisible dans le type d'origine reste du texte.
    """
    texte = _norm(texte)
    if not texte:
        return None
    if isinstance(ancienne, bool):
        return {"true": True, "false": False}.get(texte.lower(), texte)
    for genre in (dt.datetime, dt.date):
        if isinstance(ancienne, genre):
            try:
                return genre.fromisoformat(texte)
            except ValueError:
                return texte
    if isinstance(ancienne, str):
        return texte
    for conv in (int, float):
        try:
            return conv(texte)
        except ValueError:
            pass
    return texte


def valider(chemin: Path) -> None:
    """Lève `ValueError` si `sync_from_excel.py` refuserait ce classeur."""
    try:
        if detect_schema(chemin) == "single":
            parse_single(chemin)
    except ValueError:
        raise
    except Exception as e:  # fichier illisible, pas un xlsx...
        raise ValueError(f"Classeur illisible : {e}") from e


def a_change(avant: pd.DataFrame, apres: pd.DataFrame) -> bool:
    if list(apres.columns) != list(avant.columns) or not apres.index.equals(avant.index):
        return True
    return bool((apres.map(_norm) != avant).any().any())


def enregistrer(projet: Path, chemin: Path, feuille: str,
                avant: pd.DataFrame, apres: pd.DataFrame) -> Path:
    """Reporte dans l'Excel les différences entre `avant` et `apres`.

    `avant` = `lire_feuille` au moment de l'affichage ; `apres` = la sortie
    de `st.data_editor` (lignes existantes gardent leur index, lignes
    ajoutées ont un index >= len(avant), une colonne en plus dans `apres`
    est créée à droite). Lève `ValueError` si `sync_from_excel` refuse le
    résultat, `ExcelVerrouille` si le fichier ne peut pas être remplacé.
    Retourne le chemin de la sauvegarde de l'ancienne version.
    """
    wb = openpyxl.load_workbook(chemin)
    ws = wb[feuille]
    colonnes = list(avant.columns)
    for nom in apres.columns:
        if nom not in colonnes:
            colonnes.append(nom)
            ws.cell(1, len(colonnes), nom)

    derniere = len(avant) + 1
    for i, ligne in apres.iterrows():
        if i in avant.index:
            r = int(i) + 2
            for j, nom in enumerate(colonnes):
                ancien = avant.at[i, nom] if nom in avant.columns else ""
                if _norm(ligne[nom]) != ancien:
                    cellule = ws.cell(r, j + 1)
                    cellule.value = _valeur(ligne[nom], cellule.value)
        else:
            # Une cellule neuve prend le type de sa colonne (1re ligne de
            # données) : un `id` texte reste texte, un `mouse_id` reste nombre.
            derniere += 1
            for j, nom in enumerate(colonnes):
                ws.cell(derniere, j + 1).value = _valeur(ligne[nom], ws.cell(2, j + 1).value)

    # ponytail: delete_rows ne réécrit pas les références des formules
    # situées plus bas ; les templates n'en ont pas, l'Excel historique
    # multi-animaux si (feuille Subjects) — supprimer des lignes là-bas
    # reste à faire dans Excel.
    for i in sorted((i for i in avant.index if i not in apres.index), reverse=True):
        ws.delete_rows(int(i) + 2)

    sauvegardes = dossier_sauvegardes(projet)
    sauvegardes.mkdir(parents=True, exist_ok=True)
    # Fichier temporaire hors de la racine du projet : find_project_excel
    # ne doit jamais le voir comme un second classeur.
    tmp = sauvegardes / "en-cours.xlsx"
    wb.save(tmp)
    try:
        valider(tmp)
    except ValueError:
        tmp.unlink(missing_ok=True)
        raise

    sauvegarde = sauvegardes / nom_sauvegarde(chemin)
    shutil.copy2(chemin, sauvegarde)
    try:
        os.replace(tmp, chemin)
    except PermissionError as e:
        tmp.unlink(missing_ok=True)
        raise ExcelVerrouille(
            f"Impossible de remplacer `{chemin.name}` : il est probablement "
            "ouvert dans Excel. Ferme-le et réessaie."
        ) from e
    return sauvegarde


def modifie_depuis_sync(projet: Path, chemin: Path) -> bool:
    """Vrai si l'Excel est plus récent que le plus récent `metadata.yaml`."""
    metas = list(raw_dir(Path(projet)).glob("*/metadata.yaml"))
    if not metas:
        return False
    return chemin.stat().st_mtime > max(m.stat().st_mtime for m in metas)
